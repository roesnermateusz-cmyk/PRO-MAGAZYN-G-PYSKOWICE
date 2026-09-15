#!/usr/bin/env python3
"""Migracja kartoteki ze starego skoroszytu (Magazyn Zabrze) do formatu PRO-MAGAZYN.

Skrypt czyta arkusz "Dane" ze starego pliku .xlsm i zapisuje kartoteke jako
CSV w nowym ukladzie kolumn.  Dzieki temu budowanie skoroszytu jest w pelni
odtwarzalne z repozytorium, bez zaleznosci od pliku zrodlowego.

Przy okazji naprawiane sa bledy zapisane w danych przez stara wersje makr:

1. Kolumna "Wartosc" bywa zerowa mimo podanej ceny - stary formularz czytal
   kwote ze sformatowanej etykiety ("6 747,00 zl"), a IsNumeric zwracal dla
   niej False.  Wartosc jest przeliczana z wolumenu i ceny.
2. Nazwisko osoby wprowadzajacej operacje bylo doklejane do pola "Uwagi"
   jako tekst "Utworzyl: ...".  Trafia teraz do wlasnej kolumny.
3. Brak stawki za kilometr - jest wyliczana z kosztu i odleglosci
   (a gdy sie nie da, przyjmowana jest stawka domyslna).

Uruchomienie:
    python3 tools/migruj_dane.py <plik_zrodlowy.xlsm> [plik_wyjsciowy.csv]
"""

from __future__ import annotations

import csv
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path

warnings.filterwarnings("ignore")

import openpyxl  # noqa: E402

STAWKA_DOMYSLNA = 5.0

# Naglowki nowej kartoteki - kolejnosc musi odpowiadac enum kolDane w mod_Config.
NAGLOWKI = [
    "Data zaladunku",
    "Miejsce zaladunku",
    "Data operacji",
    "Dostawca",
    "Typ operacji",
    "Nr WZ",
    "Czy magazynowane",
    "Deklaracja/KZR",
    "Volumen",
    "Jednostka miary",
    "Cena zakupu/produkcji",
    "Wartosc",
    "Cena sprzedazy",
    "Produkt",
    "Rodzaj zrebki",
    "Rabanie (kto)",
    "Koszt rabania",
    "Przewoznik",
    "Nr rejestracyjny",
    "Odleglosc km",
    "Stawka zl/km",
    "Koszt transportu",
    "Odbiorca",
    "Miejsce pochodzenia",
    "Uwagi",
    "Utworzyl",
    "Data dodania wpisu",
    "ID operacji",
    "ID powiazania",
    "Status",
]

# Mapowanie: indeks kolumny w starym arkuszu (0-based) -> nazwa pola.
STARE = {
    "data_zaladunku": 0,
    "miejsce_zaladunku": 1,
    "data_operacji": 2,
    "dostawca": 3,
    "typ": 4,
    "nr_wz": 5,
    "czy_mag": 6,
    "deklaracja": 7,
    "volumen": 8,
    "jednostka": 9,
    "cena_zakupu": 10,
    "wartosc": 11,
    "cena_sprzedazy": 12,
    "produkt": 13,
    "zrebka": 14,
    "rabanie": 15,
    "koszt_rabania": 16,
    "przewoznik": 17,
    "nr_rej": 18,
    "odleglosc": 19,
    "koszt_transportu": 20,
    "odbiorca": 21,
    "uwagi": 22,
    "miejsce_pochodzenia": 23,
    "data_dodania": 24,
}

# Nazwisko konczy sie na pierwszym znaku spoza alfabetu - w danych zrodlowych
# po podpisie bywa doklejona waga ("Utworzyl: Jan Kowalski/23,78T").
WZORZEC_AUTORA = re.compile(
    r"Utworzy[lł]\s*:\s*([A-Za-zĄĆĘŁŃÓŚŹŻąćęłńóśźż][A-Za-zĄĆĘŁŃÓŚŹŻąćęłńóśźż.\- ]{2,59})",
    re.IGNORECASE,
)
# Z uwag usuwamy caly podpis wraz z ewentualnym ogonkiem po nazwisku.
WZORZEC_PODPISU = re.compile(r"Utworzy[lł]\s*:\s*[^\r\n_]*", re.IGNORECASE)


class Statystyki:
    def __init__(self) -> None:
        self.wierszy = 0
        self.naprawiona_wartosc = 0
        self.wydzielony_autor = 0
        self.wyliczona_stawka = 0
        self.pominiete = 0


def liczba(wartosc, domyslna: float = 0.0) -> float:
    if wartosc is None or wartosc == "":
        return domyslna
    try:
        return float(wartosc)
    except (TypeError, ValueError):
        return domyslna


def tekst(wartosc) -> str:
    if wartosc is None:
        return ""
    if isinstance(wartosc, datetime):
        return wartosc.strftime("%Y-%m-%d %H:%M:%S")
    return str(wartosc).replace("_x000D_", " ").replace("\r", " ").strip()


def data_iso(wartosc) -> str:
    if isinstance(wartosc, datetime):
        return wartosc.strftime("%Y-%m-%d")
    if wartosc in (None, ""):
        return ""
    return str(wartosc)


def czysc_uwagi(surowe: str) -> str:
    """Usuwa z uwag doklejony podpis 'Utworzyl: ...' i porzadkuje bialy znak."""
    bez_autora = WZORZEC_PODPISU.sub("", surowe)
    bez_autora = bez_autora.replace("_x000D_", " ")
    bez_autora = re.sub(r"[\r\n]+", " ", bez_autora)
    bez_autora = re.sub(r"\s{2,}", " ", bez_autora)
    return bez_autora.strip(" /-\t")


def migruj_wiersz(wiersz, numer: int, stat: Statystyki) -> list | None:
    def pole(nazwa):
        indeks = STARE[nazwa]
        return wiersz[indeks] if indeks < len(wiersz) else None

    typ = tekst(pole("typ")).upper()
    if not typ:
        stat.pominiete += 1
        return None

    volumen = liczba(pole("volumen"))
    cena_zakupu = liczba(pole("cena_zakupu"))
    cena_sprzedazy = liczba(pole("cena_sprzedazy"))
    wartosc = liczba(pole("wartosc"))

    # --- Naprawa 1: brakujaca wartosc operacji ---------------------------
    if wartosc == 0:
        if typ == "ZAKUP" and cena_zakupu > 0 and volumen > 0:
            wartosc = round(volumen * cena_zakupu, 2)
            stat.naprawiona_wartosc += 1
        elif typ == "SPRZEDAŻ" and cena_sprzedazy > 0 and volumen > 0:
            wartosc = round(volumen * cena_sprzedazy, 2)
            stat.naprawiona_wartosc += 1

    # --- Naprawa 2: autor wyciagniety z pola Uwagi -----------------------
    surowe_uwagi = tekst(pole("uwagi"))
    dopasowanie = WZORZEC_AUTORA.search(surowe_uwagi)
    autor = ""
    if dopasowanie:
        autor = dopasowanie.group(1).strip(" .;:/-")
        stat.wydzielony_autor += 1
    uwagi = czysc_uwagi(surowe_uwagi)

    # --- Naprawa 3: stawka za kilometr -----------------------------------
    odleglosc = liczba(pole("odleglosc"))
    koszt_transportu = liczba(pole("koszt_transportu"))
    if odleglosc > 0 and koszt_transportu > 0:
        stawka = round(koszt_transportu / odleglosc, 2)
        stat.wyliczona_stawka += 1
    elif odleglosc > 0 or koszt_transportu > 0:
        stawka = STAWKA_DOMYSLNA
    else:
        stawka = 0.0

    stat.wierszy += 1

    return [
        data_iso(pole("data_zaladunku")),
        tekst(pole("miejsce_zaladunku")),
        data_iso(pole("data_operacji")),
        tekst(pole("dostawca")),
        typ,
        tekst(pole("nr_wz")),
        tekst(pole("czy_mag")).upper(),
        tekst(pole("deklaracja")),
        volumen,
        tekst(pole("jednostka")).upper(),
        cena_zakupu,
        wartosc,
        cena_sprzedazy,
        tekst(pole("produkt")),
        tekst(pole("zrebka")),
        tekst(pole("rabanie")),
        liczba(pole("koszt_rabania")),
        tekst(pole("przewoznik")),
        tekst(pole("nr_rej")).upper(),
        odleglosc,
        stawka,
        koszt_transportu,
        tekst(pole("odbiorca")),
        tekst(pole("miejsce_pochodzenia")),
        uwagi,
        autor,
        tekst(pole("data_dodania")),
        f"OP-MIGR-{numer:05d}",
        "",
        "AKTYWNY",
    ]


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 1

    zrodlo = Path(argv[1])
    cel = Path(argv[2]) if len(argv) > 2 else Path("data/kartoteka.csv")

    if not zrodlo.exists():
        print(f"Nie znaleziono pliku: {zrodlo}")
        return 1

    skoroszyt = openpyxl.load_workbook(zrodlo, data_only=True)
    if "Dane" not in skoroszyt.sheetnames:
        print("Plik zrodlowy nie zawiera arkusza 'Dane'.")
        return 1

    arkusz = skoroszyt["Dane"]
    stat = Statystyki()
    wiersze = []

    for numer, wiersz in enumerate(arkusz.iter_rows(min_row=2, values_only=True), start=1):
        zmigrowany = migruj_wiersz(wiersz, numer, stat)
        if zmigrowany:
            wiersze.append(zmigrowany)

    # Sortowanie chronologiczne - kartoteka czyta sie wtedy jak dziennik.
    wiersze.sort(key=lambda w: (str(w[2]), str(w[26])))

    cel.parent.mkdir(parents=True, exist_ok=True)
    with cel.open("w", encoding="utf-8", newline="") as plik:
        zapis = csv.writer(plik, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        zapis.writerow(NAGLOWKI)
        zapis.writerows(wiersze)

    print(f"Zapisano: {cel}")
    print(f"  wierszy przeniesionych      : {stat.wierszy}")
    print(f"  pominietych (bez typu)      : {stat.pominiete}")
    print(f"  naprawiona kolumna Wartosc  : {stat.naprawiona_wartosc}")
    print(f"  wydzielone pole Utworzyl    : {stat.wydzielony_autor}")
    print(f"  wyliczona stawka zl/km      : {stat.wyliczona_stawka}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
