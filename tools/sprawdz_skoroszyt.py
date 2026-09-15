#!/usr/bin/env python3
"""Kontrola wygenerowanego skoroszytu PRO-MAGAZYN.

Skoroszytu nie da sie otworzyc w Excelu w srodowisku budowania, dlatego
poprawnosc pliku weryfikujemy niezaleznymi parserami: zipfile, ElementTree,
openpyxl oraz olevba.

Uruchomienie:  python3 tools/sprawdz_skoroszyt.py dist/PRO-MAGAZYN.xlsm
"""

from __future__ import annotations

import re
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

warnings.filterwarnings("ignore")

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from oletools.olevba import VBA_Parser  # noqa: E402

BLEDY: list[str] = []
OSTRZEZENIA: list[str] = []

OCZEKIWANE_ARKUSZE = {
    "START": "wsStart",
    "FORMULARZ": "wsFormularz",
    "Dane": "wsDane",
    "MAGAZYN": "wsMagazyn",
    "RAPORTY": "wsRaporty",
    "SŁOWNIK": "wsSlownik",
    "HISTORIA": "wsHistoria",
    "USTAWIENIA": "wsUstawienia",
}

OCZEKIWANE_MODULY = {
    "ThisWorkbook", "wsStart", "wsFormularz", "wsDane", "wsSlownik",
    "wsMagazyn", "wsRaporty", "wsHistoria", "wsUstawienia",
    "mod_Config", "mod_Narzedzia", "mod_Slowniki", "mod_Walidacja",
    "mod_Operacje", "mod_Formularz", "mod_Autouzupelnianie", "mod_Korekty",
    "mod_Magazyn", "mod_Raporty", "mod_ImportExport", "mod_Backup",
    "mod_Audyt", "mod_UI",
}


def sprawdz(warunek: bool, opis: str, krytyczny: bool = True) -> bool:
    if warunek:
        print(f"  [OK  ] {opis}")
    else:
        print(f"  [{'BLAD' if krytyczny else 'UWAGA'}] {opis}")
        (BLEDY if krytyczny else OSTRZEZENIA).append(opis)
    return warunek


def sprawdz_pakiet(sciezka: Path) -> dict[str, bytes]:
    print("Pakiet OPC")
    with zipfile.ZipFile(sciezka) as archiwum:
        uszkodzony = archiwum.testzip()
        sprawdz(uszkodzony is None, f"archiwum ZIP spojne ({uszkodzony or 'brak bledow'})")
        wpisy = {info.filename: archiwum.read(info.filename) for info in archiwum.infolist()}

    sprawdz("xl/vbaProject.bin" in wpisy, "projekt VBA obecny w pakiecie")

    with zipfile.ZipFile(sciezka) as archiwum:
        pierwszy = archiwum.namelist()[0]
    sprawdz(pierwszy == "[Content_Types].xml",
            f"[Content_Types].xml jest pierwsza czescia pakietu (jest: {pierwszy})")

    # Kazda czesc XML musi byc poprawnie sformowana.
    zle = []
    for nazwa, dane in wpisy.items():
        if nazwa.endswith((".xml", ".rels")):
            try:
                ET.fromstring(dane)
            except ET.ParseError as blad:
                zle.append(f"{nazwa}: {blad}")
    sprawdz(not zle, f"wszystkie czesci XML poprawne ({len(wpisy)} plikow)" if not zle else f"bledny XML: {zle[:3]}")

    typy = wpisy["[Content_Types].xml"].decode("utf-8")
    sprawdz(
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml" in typy,
        "skoroszyt zadeklarowany jako macroEnabled",
    )
    sprawdz('Extension="bin"' in typy, "typ zawartosci dla vbaProject.bin")
    sprawdz(
        typy.lstrip().startswith("<?xml") or typy.lstrip().startswith("<Types"),
        "[Content_Types].xml zaczyna sie poprawnie",
    )

    relacje = wpisy["xl/_rels/workbook.xml.rels"].decode("utf-8")
    sprawdz(
        "relationships/vbaProject" in relacje and "vbaProject.bin" in relacje,
        "relacja skoroszyt -> vbaProject.bin",
    )

    workbook_xml = wpisy["xl/workbook.xml"].decode("utf-8")
    sprawdz('codeName="ThisWorkbook"' in workbook_xml, "nazwa kodowa skoroszytu")

    return wpisy


def sprawdz_nazwy_kodowe(wpisy: dict[str, bytes]):
    print("Nazwy kodowe arkuszy")
    znalezione = set()
    for nazwa, dane in wpisy.items():
        if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", nazwa):
            tresc = dane.decode("utf-8")
            dopasowanie = re.search(r'<sheetPr[^>]*codeName="([^"]+)"', tresc)
            if dopasowanie:
                znalezione.add(dopasowanie.group(1))
            else:
                BLEDY.append(f"{nazwa}: brak codeName")
                print(f"  [BLAD] {nazwa}: brak codeName")

    oczekiwane = set(OCZEKIWANE_ARKUSZE.values())
    sprawdz(
        znalezione == oczekiwane,
        f"komplet nazw kodowych ({len(znalezione)}/{len(oczekiwane)})"
        + ("" if znalezione == oczekiwane else f" brakuje: {oczekiwane - znalezione}"),
    )


def sprawdz_vba(wpisy: dict[str, bytes]):
    print("Projekt VBA")
    vba = wpisy["xl/vbaProject.bin"]
    parser = VBA_Parser("vbaProject.bin", data=vba)
    sprawdz(parser.detect_vba_macros(), "olevba wykrywa makra")

    moduly = {}
    for _, _, nazwa_pliku, kod in parser.extract_macros():
        moduly[nazwa_pliku.rsplit(".", 1)[0]] = kod

    sprawdz(
        set(moduly) == OCZEKIWANE_MODULY,
        f"komplet modulow ({len(moduly)}/{len(OCZEKIWANE_MODULY)})"
        + ("" if set(moduly) == OCZEKIWANE_MODULY else f" roznica: {OCZEKIWANE_MODULY ^ set(moduly)}"),
    )

    # Typy modulow: arkusze i skoroszyt musza byc klasami dokumentow.
    for nazwa, kod in moduly.items():
        if nazwa.startswith("ws") or nazwa == "ThisWorkbook":
            if "Attribute VB_Base" not in kod:
                BLEDY.append(f"modul {nazwa} nie jest modulem dokumentu")
                print(f"  [BLAD] modul {nazwa} nie jest modulem dokumentu")

    sprawdz(
        all("Attribute VB_Base" in moduly[n] for n in moduly if n.startswith("ws") or n == "ThisWorkbook"),
        "moduly arkuszy maja atrybuty klasy dokumentu",
    )

    # Kluczowa logika musi byc obecna w kodzie.
    polaczony = "\n".join(moduly.values())
    for fragment, opis in (
        ('OP_TRANSPORT As String = "TRANSPORT"', "typ operacji TRANSPORT"),
        ("ZapiszTransportRownolegly", "transport jako operacja rownolegla"),
        ("PrzeliczTransport", "przeliczanie km x stawka"),
        ("kolUtworzyl", "kolumna Utworzyl"),
        ("PodpowiedzDlaPrzewoznika", "autouzupelnianie przewoznika"),
        ("OdswiezListyKartoteki", "listy rozwijane w kolumnach kartoteki"),
        ("UstawListeRozwijana", "zakladanie list rozwijanych"),
        ("KorygujWiersz", "korekty (storno)"),
        ("UtworzKopieZapasowa", "kopie zapasowe"),
        ("ZapisDoHistorii", "dziennik zdarzen"),
        ("RaportRoczny", "zestawienie roczne"),
        ("ImportujKartoteke", "import CSV"),
    ):
        sprawdz(fragment in polaczony, f"kod zawiera: {opis}")

    # Polskie znaki musza przetrwac kodowanie cp1250.
    sprawdz("SPRZEDAŻ" in polaczony, "polskie znaki w kodzie VBA")


def sprawdz_arkusze(sciezka: Path):
    print("Zawartosc skoroszytu")
    skoroszyt = openpyxl.load_workbook(sciezka, keep_vba=True)
    sprawdz(
        set(skoroszyt.sheetnames) == set(OCZEKIWANE_ARKUSZE),
        f"komplet arkuszy: {skoroszyt.sheetnames}",
    )

    dane = skoroszyt["Dane"]
    sprawdz("Tabela_dane" in dane.tables, "tabela Tabela_dane istnieje")
    if "Tabela_dane" in dane.tables:
        tabela = dane.tables["Tabela_dane"]
        sprawdz(len(tabela.tableColumns) == 43, f"tabela ma 43 kolumny (jest {len(tabela.tableColumns)})")

    naglowki = [komorka.value for komorka in dane[1]]
    for wymagana in ("Typ operacji", "Utworzył", "Stawka zł/km", "Koszt transportu", "Status", "ID operacji"):
        sprawdz(wymagana in naglowki, f"kolumna kartoteki: {wymagana}")

    # Nazwane zakresy slownikow i przelicznikow.
    nazwy = set(skoroszyt.defined_names)
    for wymagana in ("sl_PRODUKT", "sl_PRZEWOZNIK", "sl_OPERATOR", "sl_TYP_OPERACJI",
                     "Przelicznik_MP_tona", "Przelicznik_M3_MP", "Przelicznik_GJ_tona"):
        sprawdz(wymagana in nazwy, f"nazwany zakres: {wymagana}")

    # Kazda lista rozwijana musi wskazywac istniejacy nazwany zakres.
    brakujace = set()
    for arkusz in skoroszyt.worksheets:
        for sprawdzanie in arkusz.data_validations.dataValidation:
            formula = (sprawdzanie.formula1 or "").lstrip("=")
            if formula.startswith("sl_") and formula not in nazwy:
                brakujace.add(formula)
    sprawdz(not brakujace, f"listy rozwijane wskazuja istniejace zakresy ({brakujace or 'wszystkie OK'})")

    # Slownik operatorow musi zawierac osoby z kartoteki.
    slownik = skoroszyt["SŁOWNIK"]
    naglowki_slownika = [komorka.value for komorka in slownik[1]]
    sprawdz("OPERATOR" in naglowki_slownika, "slownik OPERATOR (pole Utworzyl)")
    if "OPERATOR" in naglowki_slownika:
        kolumna = naglowki_slownika.index("OPERATOR") + 1
        osoby = [
            slownik.cell(row=w, column=kolumna).value
            for w in range(2, 40)
            if slownik.cell(row=w, column=kolumna).value
        ]
        sprawdz(len(osoby) >= 2, f"slownik operatorow wypelniony: {osoby}")

    formularz = skoroszyt["FORMULARZ"]
    klucze = {
        formularz.cell(row=w, column=1).value
        for w in range(1, 80)
        if formularz.cell(row=w, column=1).value
    }
    for wymagany in ("TYP", "UTWORZYL", "PRZEWOZNIK", "NR_REJ", "ODLEGLOSC",
                     "STAWKA_KM", "KOSZT_TRANSPORTU", "ROWN_TRANSPORT",
                     "ROWN_PRODUKCJA", "ROWN_SPRZEDAZ"):
        sprawdz(wymagany in klucze, f"pole formularza: {wymagany}")

    ustawienia = skoroszyt["USTAWIENIA"]
    klucze_ustawien = {
        ustawienia.cell(row=w, column=2).value: ustawienia.cell(row=w, column=3).value
        for w in range(2, 30)
        if ustawienia.cell(row=w, column=2).value
    }
    sprawdz(
        klucze_ustawien.get("STAWKA_TRANSPORT_KM") == 5,
        f"stawka transportu = 5 zl/km (jest {klucze_ustawien.get('STAWKA_TRANSPORT_KM')})",
    )

    return skoroszyt


# Funkcje dozwolone w formulach kartoteki - literowka w nazwie daje w Excelu #NAZWA?.
FUNKCJE_EXCELA = {
    "IF", "OR", "AND", "NOT", "ISNUMBER", "SEARCH", "UPPER", "LOWER", "TEXT",
    "YEAR", "MONTH", "DAY", "ROUNDUP", "ROUND", "ROUNDDOWN", "MAX", "MIN",
    "SUM", "COUNTA", "OFFSET", "IFERROR", "ABS", "INT", "LEN", "TRIM", "N",
}


def sprawdz_formuly(skoroszyt):
    """Kontroluje formuly kartoteki: nazwy funkcji i istnienie nazwanych zakresow."""
    print("Formuly kartoteki")
    dane = skoroszyt["Dane"]
    nazwy = set(skoroszyt.defined_names)
    naglowki = [komorka.value for komorka in dane[1]]

    nieznane_funkcje = set()
    nieznane_nazwy = set()
    policzone = 0

    for numer in range(2, min(dane.max_row, 200) + 1):
        for kolumna in range(31, len(naglowki) + 1):
            formula = dane.cell(row=numer, column=kolumna).value
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            policzone += 1
            # Literaly tekstowe nie sa identyfikatorami - usuwamy je przed analiza.
            bez_tekstow = re.sub(r'"[^"]*"', '""', formula)
            for token in re.findall(r"\b([A-Za-z_][A-Za-z0-9_.]*)\s*\(", bez_tekstow):
                if token.upper() not in FUNKCJE_EXCELA:
                    nieznane_funkcje.add(token)
            # Identyfikatory nie bedace funkcjami ani adresami komorek
            # musza byc nazwanymi zakresami.
            for token in re.findall(r"\b([A-Za-z_][A-Za-z0-9_]{2,})\b(?!\s*\()", bez_tekstow):
                if token.upper() in FUNKCJE_EXCELA:
                    continue
                if re.fullmatch(r"[A-Z]{1,3}\d+", token):
                    continue
                if token not in nazwy:
                    nieznane_nazwy.add(token)

    sprawdz(policzone > 0, f"kolumny wyliczane zawieraja formuly ({policzone} sprawdzonych)")
    sprawdz(not nieznane_funkcje, f"wszystkie funkcje rozpoznane ({nieznane_funkcje or 'OK'})")
    sprawdz(not nieznane_nazwy, f"wszystkie nazwy zakresow istnieja ({nieznane_nazwy or 'OK'})")

    # Odwolanie do kolumny statusu musi trafiac w kolumne AD (30).
    przyklad = dane.cell(row=2, column=34).value or ""
    sprawdz("AD2" in przyklad, "formula ruchu magazynowego czyta kolumne Status (AD)")


def sprawdz_dane_kartoteki(skoroszyt):
    print("Poprawnosc danych kartoteki")
    dane = skoroszyt["Dane"]
    naglowki = [komorka.value for komorka in dane[1]]
    idx = {nazwa: numer + 1 for numer, nazwa in enumerate(naglowki)}

    wierszy = 0
    bez_wartosci = 0
    bez_autora = 0
    bez_statusu = 0

    for numer in range(2, dane.max_row + 1):
        typ = dane.cell(row=numer, column=idx["Typ operacji"]).value
        if not typ:
            continue
        wierszy += 1
        wartosc = dane.cell(row=numer, column=idx["Wartość"]).value or 0
        cena = dane.cell(row=numer, column=idx["Cena zakupu/produkcji"]).value or 0
        volumen = dane.cell(row=numer, column=idx["Volumen"]).value or 0
        if typ == "ZAKUP" and cena and volumen and not wartosc:
            bez_wartosci += 1
        if not dane.cell(row=numer, column=idx["Utworzył"]).value:
            bez_autora += 1
        if not dane.cell(row=numer, column=idx["Status"]).value:
            bez_statusu += 1

    sprawdz(wierszy > 0, f"kartoteka zawiera dane ({wierszy} wierszy)")
    sprawdz(bez_wartosci == 0, f"brak wierszy ZAKUP z ceną i zerową wartością (jest {bez_wartosci})")
    sprawdz(bez_statusu == 0, f"kazdy wiersz ma status (bez statusu: {bez_statusu})")
    sprawdz(
        bez_autora < wierszy,
        f"pole Utworzyl wypelnione dla {wierszy - bez_autora}/{wierszy} wierszy",
        krytyczny=False,
    )


def main(argv: list[str]) -> int:
    sciezka = Path(argv[1] if len(argv) > 1 else "dist/PRO-MAGAZYN.xlsm")
    if not sciezka.exists():
        print(f"Nie znaleziono pliku: {sciezka}")
        return 1

    print(f"Kontrola pliku: {sciezka}  ({sciezka.stat().st_size/1024:.0f} KB)\n")

    wpisy = sprawdz_pakiet(sciezka)
    sprawdz_nazwy_kodowe(wpisy)
    sprawdz_vba(wpisy)
    skoroszyt = sprawdz_arkusze(sciezka)
    sprawdz_formuly(skoroszyt)
    sprawdz_dane_kartoteki(skoroszyt)

    print()
    if OSTRZEZENIA:
        print(f"Ostrzezenia: {len(OSTRZEZENIA)}")
    if BLEDY:
        print(f"NIEPOWODZENIE - bledow: {len(BLEDY)}")
        for blad in BLEDY:
            print("  -", blad)
        return 1
    print("Skoroszyt przeszedl wszystkie kontrole.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
