"""Odczyt pliku danych (.xlsx / .xlsm) do modelu aplikacji."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..model.operacja import Operacja, Status, TypOperacji, ZdarzenieAudytu
from ..model.slowniki import Slowniki
from ..model.ustawienia import Ustawienia
from . import uklad


class BladOdczytu(Exception):
    """Plik danych jest nieczytelny albo ma nieznany układ."""


def _liczba(wartosc, domyslna: float = 0.0) -> float:
    if wartosc is None or wartosc == "":
        return domyslna
    if isinstance(wartosc, (int, float)):
        return float(wartosc)
    tekst = str(wartosc).strip().replace(" ", "").replace("\xa0", "")
    tekst = "".join(znak for znak in tekst if znak.isdigit() or znak in ",.-")
    if tekst.count(",") and tekst.count("."):
        tekst = tekst.replace(".", "")
    tekst = tekst.replace(",", ".")
    try:
        return float(tekst)
    except ValueError:
        return domyslna


def _data(wartosc) -> date | None:
    if isinstance(wartosc, datetime):
        return wartosc.date()
    if isinstance(wartosc, date):
        return wartosc
    if not wartosc:
        return None
    for wzorzec in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(str(wartosc).strip(), wzorzec).date()
        except ValueError:
            continue
    return None


def _data_czasu(wartosc) -> datetime | None:
    if isinstance(wartosc, datetime):
        return wartosc
    if isinstance(wartosc, date):
        return datetime.combine(wartosc, datetime.min.time())
    if not wartosc:
        return None
    for wzorzec in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y"):
        try:
            return datetime.strptime(str(wartosc).strip(), wzorzec)
        except ValueError:
            continue
    return None


def _tekst(wartosc) -> str:
    if wartosc is None:
        return ""
    if isinstance(wartosc, float) and wartosc.is_integer():
        return str(int(wartosc))
    return str(wartosc).strip()


def wiersz_na_operacje(wiersz: tuple) -> Operacja | None:
    """Zamienia wiersz arkusza Dane na operację. Zwraca None dla pustych wierszy."""
    def pole(numer: int):
        return wiersz[numer - 1] if numer - 1 < len(wiersz) else None

    dane: dict[str, object] = {}
    for numer, nazwa_pola in enumerate(uklad.POLA_OPERACJI, start=1):
        wartosc = pole(numer)

        if nazwa_pola == "typ":
            typ = TypOperacji.z_tekstu(_tekst(wartosc))
            if typ is None:
                return None
            dane[nazwa_pola] = typ
        elif nazwa_pola == "status":
            tekst = _tekst(wartosc).upper()
            dane[nazwa_pola] = next(
                (s for s in Status if s.value == tekst), Status.AKTYWNY
            )
        elif nazwa_pola == "czy_magazynowane":
            dane[nazwa_pola] = _tekst(wartosc).upper() != "NIE"
        elif nazwa_pola in uklad.POLA_LICZBOWE:
            dane[nazwa_pola] = _liczba(wartosc)
        elif nazwa_pola in uklad.POLA_DATY:
            dane[nazwa_pola] = _data(wartosc)
        elif nazwa_pola in uklad.POLA_DATY_CZASU:
            dane[nazwa_pola] = _data_czasu(wartosc)
        else:
            dane[nazwa_pola] = _tekst(wartosc)

    return Operacja(**dane)


def wczytaj(sciezka: Path) -> tuple[list[Operacja], Slowniki, Ustawienia, list[ZdarzenieAudytu]]:
    """Wczytuje komplet danych z pliku Excela."""
    if not sciezka.exists():
        raise BladOdczytu(f"Plik nie istnieje: {sciezka}")

    try:
        skoroszyt = load_workbook(sciezka, data_only=True, read_only=False)
    except Exception as blad:  # noqa: BLE001 - komunikat trafia do użytkownika
        raise BladOdczytu(f"Nie udało się otworzyć pliku: {blad}") from blad

    if "Dane" not in skoroszyt.sheetnames:
        raise BladOdczytu(
            "Plik nie zawiera arkusza „Dane” - to nie jest plik danych PRO-MAGAZYN."
        )

    operacje: list[Operacja] = []
    for wiersz in skoroszyt["Dane"].iter_rows(min_row=2, values_only=True):
        operacja = wiersz_na_operacje(wiersz)
        if operacja is not None:
            operacje.append(operacja)

    ustawienia = _wczytaj_ustawienia(skoroszyt)
    zdarzenia = _wczytaj_historie(skoroszyt)

    # Słowniki bierzemy z arkusza i uzupełniamy o wartości z kartoteki,
    # żeby listy obejmowały także pozycje dopisane ręcznie w Excelu.
    slowniki = Slowniki()
    if "SŁOWNIK" in skoroszyt.sheetnames:
        _wczytaj_slowniki(skoroszyt["SŁOWNIK"], slowniki)
    for operacja in operacje:
        slowniki.uzupelnij_z_operacji(operacja)
    slowniki.posortuj()

    skoroszyt.close()
    return operacje, slowniki, ustawienia, zdarzenia


def _wczytaj_ustawienia(skoroszyt) -> Ustawienia:
    ustawienia = Ustawienia()
    if "USTAWIENIA" not in skoroszyt.sheetnames:
        return ustawienia

    arkusz = skoroszyt["USTAWIENIA"]
    for wiersz in arkusz.iter_rows(min_row=2, min_col=2, max_col=3, values_only=True):
        klucz = _tekst(wiersz[0]).upper()
        if klucz:
            ustawienia.ustaw(klucz, wiersz[1] if wiersz[1] is not None else "")
    return ustawienia


def _wczytaj_historie(skoroszyt) -> list[ZdarzenieAudytu]:
    if "HISTORIA" not in skoroszyt.sheetnames:
        return []

    zdarzenia: list[ZdarzenieAudytu] = []
    for wiersz in skoroszyt["HISTORIA"].iter_rows(min_row=2, values_only=True):
        data = _data_czasu(wiersz[1] if len(wiersz) > 1 else None)
        if data is None:
            continue
        zdarzenia.append(
            ZdarzenieAudytu(
                data=data,
                uzytkownik=_tekst(wiersz[2] if len(wiersz) > 2 else ""),
                zdarzenie=_tekst(wiersz[3] if len(wiersz) > 3 else ""),
                typ_operacji=_tekst(wiersz[4] if len(wiersz) > 4 else ""),
                id_operacji=_tekst(wiersz[5] if len(wiersz) > 5 else ""),
                opis=_tekst(wiersz[6] if len(wiersz) > 6 else ""),
            )
        )
    return zdarzenia


def _wczytaj_slowniki(arkusz, slowniki: Slowniki) -> None:
    naglowki: dict[int, str] = {}
    for numer_kolumny, komorka in enumerate(arkusz[1], start=1):
        klucz = _tekst(komorka.value).upper()
        if klucz:
            naglowki[numer_kolumny] = klucz

    for numer_kolumny, klucz in naglowki.items():
        for wiersz in arkusz.iter_rows(
            min_row=2, min_col=numer_kolumny, max_col=numer_kolumny, values_only=True
        ):
            wartosc = _tekst(wiersz[0])
            if wartosc:
                slowniki.dodaj(klucz, wartosc)
