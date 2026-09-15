"""Budowa skoroszytu Excela z danych aplikacji.

To jedyne miejsce, w którym powstaje plik .xlsx. Korzysta z niego zarówno
aplikacja desktopowa (zapis kartoteki), jak i generator wersji z makrami
(build/build_workbook.py), dzięki czemu oba formaty mają identyczny układ.
"""

from __future__ import annotations

from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..model.operacja import Operacja, Status, TypOperacji, ZdarzenieAudytu
from ..model.slowniki import Slowniki
from ..model.ustawienia import DOMYSLNE, Ustawienia
from . import uklad

GRANAT = "1F3A5F"
ZIELONY = "10845A"
SZARY = "5A626E"
BIALY = "FFFFFF"

CIENKA = Side(style="thin", color="C8D0DA")
RAMKA = Border(left=CIENKA, right=CIENKA, top=CIENKA, bottom=CIENKA)

ARKUSZE = {
    "START": "wsStart",
    "FORMULARZ": "wsFormularz",
    "Dane": "wsDane",
    "MAGAZYN": "wsMagazyn",
    "RAPORTY": "wsRaporty",
    "SŁOWNIK": "wsSlownik",
    "HISTORIA": "wsHistoria",
    "USTAWIENIA": "wsUstawienia",
}


def _naglowek(komorka, tlo=GRANAT, rozmiar=11):
    komorka.font = Font(bold=True, color=BIALY, size=rozmiar, name="Segoe UI")
    komorka.fill = PatternFill("solid", fgColor=tlo)
    komorka.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    komorka.border = RAMKA


def _tytul(arkusz, komorka, tekst, rozmiar=18):
    arkusz[komorka] = tekst
    arkusz[komorka].font = Font(bold=True, size=rozmiar, color=GRANAT, name="Segoe UI")


# ----------------------------------------------------------------------
#  Konwersja modelu na wiersz arkusza
# ----------------------------------------------------------------------


def operacja_na_wiersz(op: Operacja) -> list:
    """Zamienia operację na listę wartości kolumn 1-30."""
    wiersz = []
    for nazwa_pola in uklad.POLA_OPERACJI:
        wartosc = getattr(op, nazwa_pola)

        if nazwa_pola == "typ":
            wartosc = wartosc.value if isinstance(wartosc, TypOperacji) else str(wartosc or "")
        elif nazwa_pola == "status":
            wartosc = wartosc.value if isinstance(wartosc, Status) else str(wartosc or "")
        elif nazwa_pola == "czy_magazynowane":
            wartosc = "TAK" if wartosc else "NIE"
        elif nazwa_pola in uklad.POLA_LICZBOWE:
            wartosc = float(wartosc or 0)
        elif wartosc is None:
            wartosc = None
        elif not isinstance(wartosc, (date, datetime, int, float)):
            wartosc = str(wartosc)

        wiersz.append(wartosc)
    return wiersz


# ----------------------------------------------------------------------
#  Arkusze
# ----------------------------------------------------------------------


def _zbuduj_start(arkusz, ust: Ustawienia):
    arkusz.sheet_view.showGridLines = False
    for kolumna, szerokosc in zip("ABCDEFGHI", [2, 22, 18, 18, 18, 22, 18, 18, 4]):
        arkusz.column_dimensions[kolumna].width = szerokosc

    _tytul(arkusz, "B2", "PRO-MAGAZYN", 26)
    arkusz["B3"] = "System ewidencji magazynowej - biomasa i drewno"
    arkusz["B3"].font = Font(size=12, color=SZARY, name="Segoe UI")
    arkusz["B4"] = f"{ust.tekst('FIRMA', 'ResInvest Commodities')}  |  wersja 2.1.0"
    arkusz["B4"].font = Font(size=10, italic=True, color=SZARY, name="Segoe UI")

    for kolumna, etykieta in (
        ("B", "Operacji w kartotece"),
        ("D", "Zdarzeń w dzienniku"),
        ("F", "Ostatnie odświeżenie"),
    ):
        arkusz[f"{kolumna}6"] = etykieta
        arkusz[f"{kolumna}6"].font = Font(size=9, bold=True, color=SZARY, name="Segoe UI")

    for komorka in ("C7", "E7", "G7"):
        arkusz[komorka].font = Font(size=14, bold=True, color=GRANAT, name="Segoe UI")
    arkusz["G7"].number_format = "dd.mm.yyyy hh:mm"

    arkusz["B8"] = "Przyciski poniżej uruchamiają moduły systemu (wersja z makrami)."
    arkusz["B8"].font = Font(size=9, italic=True, color=SZARY, name="Segoe UI")

    for wiersz in range(10, 23):
        arkusz.row_dimensions[wiersz].height = 23

    arkusz["B24"] = (
        "Skróty:  Ctrl+Shift+N nowa operacja   |   Ctrl+Shift+Z zapisz   |   "
        "Ctrl+Shift+M magazyn   |   Ctrl+Shift+R raporty"
    )
    arkusz["B24"].font = Font(size=9, color=SZARY, name="Segoe UI")


def _zbuduj_formularz(arkusz):
    arkusz.sheet_view.showGridLines = False
    arkusz.column_dimensions["A"].hidden = True
    arkusz.column_dimensions["A"].width = 2
    for kolumna, szerokosc in (("B", 34), ("C", 30), ("D", 48), ("E", 3), ("F", 16), ("G", 16)):
        arkusz.column_dimensions[kolumna].width = szerokosc

    _tytul(arkusz, "B1", "NOWA OPERACJA MAGAZYNOWA", 16)
    arkusz["B2"] = "Pola oznaczone * są obowiązkowe. Wartości wybieraj z list rozwijanych."
    arkusz["B2"].font = Font(size=9, italic=True, color=SZARY, name="Segoe UI")

    wiersz = 4
    pola: dict[str, int] = {}

    for pozycja in uklad.FORMULARZ:
        if pozycja[0] == "blank":
            arkusz.row_dimensions[wiersz].height = 6
            wiersz += 1
            continue

        if pozycja[0] == "section":
            arkusz.cell(row=wiersz, column=2, value=pozycja[1])
            for kolumna in range(2, 5):
                komorka = arkusz.cell(row=wiersz, column=kolumna)
                komorka.fill = PatternFill("solid", fgColor=GRANAT)
                komorka.font = Font(bold=True, color=BIALY, size=10, name="Segoe UI")
            arkusz.row_dimensions[wiersz].height = 20
            wiersz += 1
            continue

        _, klucz, etykieta, podpowiedz, format_liczb = pozycja
        arkusz.cell(row=wiersz, column=1, value=klucz)

        komorka = arkusz.cell(row=wiersz, column=2, value=etykieta)
        komorka.font = Font(size=10, bold=etykieta.endswith("*"), name="Segoe UI")
        komorka.alignment = Alignment(vertical="center")

        wartosc = arkusz.cell(row=wiersz, column=3)
        wartosc.fill = PatternFill("solid", fgColor=BIALY)
        wartosc.border = RAMKA
        wartosc.number_format = format_liczb
        wartosc.font = Font(size=11, name="Segoe UI")
        wartosc.alignment = Alignment(vertical="center")

        arkusz.cell(row=wiersz, column=4, value=podpowiedz).font = Font(
            size=9, italic=True, color=SZARY, name="Segoe UI"
        )

        arkusz.row_dimensions[wiersz].height = 18
        pola[klucz] = wiersz
        wiersz += 1

    for klucz, slownik in uklad.LISTY_FORMULARZA.items():
        if klucz not in pola:
            continue
        sprawdzanie = DataValidation(
            type="list",
            formula1=f"sl_{slownik}",
            allow_blank=True,
            showDropDown=False,
            errorStyle="information",
            error="Ta pozycja nie występuje jeszcze w słowniku - zostanie do niego dopisana.",
            errorTitle="Wartość spoza słownika",
        )
        arkusz.add_data_validation(sprawdzanie)
        sprawdzanie.add(arkusz.cell(row=pola[klucz], column=3))

    arkusz.cell(row=pola["TYP"], column=3, value="ZAKUP")
    arkusz.cell(row=pola["CZY_MAG"], column=3, value="TAK")
    arkusz.cell(row=pola["STAWKA_KM"], column=3, value=5)
    for klucz in ("ROWN_PRODUKCJA", "ROWN_SPRZEDAZ", "ROWN_TRANSPORT"):
        arkusz.cell(row=pola[klucz], column=3, value="NIE")


def _zbuduj_dane(arkusz, operacje: list[Operacja]) -> int:
    arkusz.freeze_panes = "A2"

    wszystkie = uklad.KOLUMNY_DANE + [nazwa for nazwa, _, _ in uklad.KOLUMNY_WYLICZANE]
    for numer, nazwa in enumerate(wszystkie, start=1):
        _naglowek(arkusz.cell(row=1, column=numer, value=nazwa), rozmiar=9)
    arkusz.row_dimensions[1].height = 34

    for przesuniecie, op in enumerate(operacje):
        numer_wiersza = przesuniecie + 2
        for numer_kolumny, wartosc in enumerate(operacja_na_wiersz(op), start=1):
            arkusz.cell(row=numer_wiersza, column=numer_kolumny, value=wartosc)

    ostatni = max(len(operacje) + 1, 2)

    for przesuniecie, (_, szablon, format_liczb) in enumerate(uklad.KOLUMNY_WYLICZANE):
        numer_kolumny = len(uklad.KOLUMNY_DANE) + 1 + przesuniecie
        for numer_wiersza in range(2, ostatni + 1):
            komorka = arkusz.cell(row=numer_wiersza, column=numer_kolumny)
            komorka.value = szablon.format(r=numer_wiersza)
            komorka.number_format = format_liczb

    formaty = {
        1: "dd.mm.yyyy", 3: "dd.mm.yyyy", 9: "# ##0.000", 11: "# ##0.00 zł",
        12: "# ##0.00 zł", 13: "# ##0.00 zł", 17: "# ##0.00 zł", 20: "# ##0.00",
        21: "# ##0.00 zł", 22: "# ##0.00 zł", 27: "dd.mm.yyyy hh:mm",
    }
    for numer_kolumny, format_liczb in formaty.items():
        for numer_wiersza in range(2, ostatni + 1):
            arkusz.cell(row=numer_wiersza, column=numer_kolumny).number_format = format_liczb

    szerokosci = {
        1: 13, 2: 20, 3: 13, 4: 28, 5: 13, 6: 11, 7: 10, 8: 13, 9: 11, 10: 8,
        11: 14, 12: 15, 13: 14, 14: 30, 15: 9, 16: 16, 17: 12, 18: 24, 19: 14,
        20: 11, 21: 11, 22: 14, 23: 24, 24: 20, 25: 34, 26: 20, 27: 17, 28: 20,
        29: 20, 30: 14,
    }
    for numer_kolumny, szerokosc in szerokosci.items():
        arkusz.column_dimensions[get_column_letter(numer_kolumny)].width = szerokosc
    for numer_kolumny in range(len(uklad.KOLUMNY_DANE) + 1, len(wszystkie) + 1):
        arkusz.column_dimensions[get_column_letter(numer_kolumny)].width = 15

    ostatnia_litera = get_column_letter(len(wszystkie))
    tabela = Table(displayName="Tabela_dane", ref=f"A1:{ostatnia_litera}{ostatni}")
    tabela.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    arkusz.add_table(tabela)

    for numer_kolumny, slownik in uklad.LISTY_KARTOTEKI.items():
        litera = get_column_letter(numer_kolumny)
        sprawdzanie = DataValidation(
            type="list",
            formula1=f"sl_{slownik}",
            allow_blank=True,
            showDropDown=False,
            errorStyle="information",
            error="Wartość spoza słownika - zostanie dopisana przy najbliższym odświeżeniu.",
            errorTitle="Wartość spoza słownika",
        )
        arkusz.add_data_validation(sprawdzanie)
        sprawdzanie.add(f"{litera}2:{litera}{max(ostatni, 5000)}")

    arkusz.conditional_formatting.add(
        f"A2:{ostatnia_litera}{ostatni}",
        CellIsRule(
            operator="equal",
            formula=['"SKORYGOWANY"'],
            stopIfTrue=False,
            fill=PatternFill("solid", fgColor="FFE4E4"),
        ),
    )
    return ostatni


def _zbuduj_magazyn(arkusz):
    arkusz.sheet_view.showGridLines = False
    for kolumna, szerokosc in zip("ABCDEFGH", [2, 24, 34, 14, 14, 14, 16, 18]):
        arkusz.column_dimensions[kolumna].width = szerokosc

    _tytul(arkusz, "B1", "STAN MAGAZYNOWY", 16)
    arkusz["F2"] = "Przeliczono:"
    arkusz["F2"].font = Font(size=9, color=SZARY, name="Segoe UI")
    arkusz["G2"].number_format = "dd.mm.yyyy hh:mm"

    for kolumna, etykieta in (("B", "RAZEM [MP]"), ("D", "RAZEM [tony]"), ("F", "RAZEM [GJ]")):
        arkusz[f"{kolumna}4"] = etykieta
        arkusz[f"{kolumna}4"].font = Font(size=9, bold=True, color=SZARY, name="Segoe UI")
    for komorka in ("C4", "E4", "G4"):
        arkusz[komorka].font = Font(size=13, bold=True, color=ZIELONY, name="Segoe UI")
        arkusz[komorka].number_format = "# ##0.000"

    for przesuniecie, nazwa in enumerate(uklad.NAGLOWKI_MAGAZYN):
        _naglowek(arkusz.cell(row=6, column=2 + przesuniecie, value=nazwa), rozmiar=10)
    arkusz.row_dimensions[6].height = 26
    arkusz.freeze_panes = "A7"


def _zbuduj_raporty(arkusz):
    arkusz.sheet_view.showGridLines = False
    szerokosci = [2, 14, 16, 14, 16, 8, 16, 22, 22, 18, 16, 12, 18, 14, 3, 14, 14, 14]
    for kolumna, szerokosc in zip("ABCDEFGHIJKLMNOPQR", szerokosci):
        arkusz.column_dimensions[kolumna].width = szerokosc

    _tytul(arkusz, "B1", "RAPORTY I ZESTAWIENIA", 16)

    for komorka, tekst in (
        ("B3", "Rok"), ("D3", "Miesiąc"), ("F3", "Typ operacji"),
        ("B4", "Produkt"), ("D4", "Kontrahent"), ("F4", "Lokalizacja"),
    ):
        arkusz[komorka] = tekst
        arkusz[komorka].font = Font(size=9, bold=True, color=SZARY, name="Segoe UI")

    for komorka in ("C3", "E3", "G3", "C4", "E4", "G4"):
        arkusz[komorka].fill = PatternFill("solid", fgColor=BIALY)
        arkusz[komorka].border = RAMKA
        arkusz[komorka].font = Font(size=10, name="Segoe UI")

    arkusz["C3"] = datetime.now().year
    for komorka in ("E3", "G3", "C4", "E4", "G4"):
        arkusz[komorka] = "WSZYSTKIE"

    miesiace = DataValidation(
        type="list",
        formula1='"WSZYSTKIE,STYCZEŃ,LUTY,MARZEC,KWIECIEŃ,MAJ,CZERWIEC,LIPIEC,'
        'SIERPIEŃ,WRZESIEŃ,PAŹDZIERNIK,LISTOPAD,GRUDZIEŃ"',
        allow_blank=True,
    )
    arkusz.add_data_validation(miesiace)
    miesiace.add(arkusz["E3"])

    for komorka, tekst in (
        ("B7", "Operacji"), ("D7", "Wolumen [MP]"), ("F7", "Wolumen [t]"),
        ("H7", "Zakup"), ("J7", "Sprzedaż"), ("L7", "Transport"),
        ("J8", "Marża"), ("L8", "Rąbanie"),
    ):
        arkusz[komorka] = tekst
        arkusz[komorka].font = Font(size=9, bold=True, color=SZARY, name="Segoe UI")

    for komorka, format_liczb in (
        ("C7", "# ##0"), ("E7", "# ##0.000"), ("G7", "# ##0.000"),
        ("I7", "# ##0.00 zł"), ("K7", "# ##0.00 zł"), ("M7", "# ##0.00 zł"),
        ("K8", "# ##0.00 zł"), ("M8", "# ##0.00 zł"),
    ):
        arkusz[komorka].font = Font(size=11, bold=True, color=GRANAT, name="Segoe UI")
        arkusz[komorka].number_format = format_liczb

    for przesuniecie, nazwa in enumerate(uklad.NAGLOWKI_RAPORT):
        _naglowek(arkusz.cell(row=10, column=2 + przesuniecie, value=nazwa), rozmiar=9)
    arkusz.row_dimensions[10].height = 26
    arkusz.freeze_panes = "A11"


def _zbuduj_slownik(arkusz, slowniki: Slowniki):
    arkusz.freeze_panes = "A2"
    for numer_kolumny, (klucz, _) in enumerate(uklad.SLOWNIKI, start=1):
        _naglowek(arkusz.cell(row=1, column=numer_kolumny, value=klucz), rozmiar=9)
        arkusz.column_dimensions[get_column_letter(numer_kolumny)].width = 26
        for przesuniecie, pozycja in enumerate(slowniki.lista(klucz)):
            komorka = arkusz.cell(row=2 + przesuniecie, column=numer_kolumny, value=pozycja)
            komorka.font = Font(size=10, name="Segoe UI")
    arkusz.row_dimensions[1].height = 28


def _zbuduj_historie(arkusz, zdarzenia: list[ZdarzenieAudytu]):
    arkusz.freeze_panes = "A2"
    for przesuniecie, nazwa in enumerate(uklad.NAGLOWKI_HISTORIA, start=1):
        _naglowek(arkusz.cell(row=1, column=przesuniecie, value=nazwa), rozmiar=10)
    for kolumna, szerokosc in zip("ABCDEFG", [8, 20, 24, 20, 16, 26, 80]):
        arkusz.column_dimensions[kolumna].width = szerokosc
    arkusz.row_dimensions[1].height = 26

    for przesuniecie, wpis in enumerate(zdarzenia):
        numer_wiersza = 2 + przesuniecie
        arkusz.cell(row=numer_wiersza, column=1, value=przesuniecie + 1)
        komorka = arkusz.cell(row=numer_wiersza, column=2, value=wpis.data)
        komorka.number_format = "dd.mm.yyyy hh:mm:ss"
        arkusz.cell(row=numer_wiersza, column=3, value=wpis.uzytkownik)
        arkusz.cell(row=numer_wiersza, column=4, value=wpis.zdarzenie)
        arkusz.cell(row=numer_wiersza, column=5, value=wpis.typ_operacji)
        arkusz.cell(row=numer_wiersza, column=6, value=wpis.id_operacji)
        arkusz.cell(row=numer_wiersza, column=7, value=wpis.opis)


def _zbuduj_ustawienia(arkusz, ust: Ustawienia) -> dict[str, int]:
    arkusz.sheet_view.showGridLines = False
    for kolumna, szerokosc in zip("ABCD", [2, 32, 26, 66]):
        arkusz.column_dimensions[kolumna].width = szerokosc

    for przesuniecie, nazwa in enumerate(("Klucz", "Wartość", "Opis")):
        _naglowek(arkusz.cell(row=1, column=2 + przesuniecie, value=nazwa), rozmiar=10)

    wiersze: dict[str, int] = {}
    for przesuniecie, (klucz, _, opis) in enumerate(DOMYSLNE):
        numer_wiersza = 2 + przesuniecie
        arkusz.cell(row=numer_wiersza, column=2, value=klucz).font = Font(
            size=10, bold=True, name="Consolas"
        )
        wartosc = ust.get(klucz, "")
        komorka = arkusz.cell(row=numer_wiersza, column=3, value=wartosc if wartosc != "" else None)
        komorka.fill = PatternFill("solid", fgColor=BIALY)
        komorka.border = RAMKA
        komorka.font = Font(size=11, name="Segoe UI")
        arkusz.cell(row=numer_wiersza, column=4, value=opis).font = Font(
            size=9, italic=True, color=SZARY, name="Segoe UI"
        )
        wiersze[klucz] = numer_wiersza

    return wiersze


def _dodaj_nazwy(skoroszyt, wiersze_ustawien: dict[str, int]):
    for numer_kolumny, (klucz, _) in enumerate(uklad.SLOWNIKI, start=1):
        litera = get_column_letter(numer_kolumny)
        odwolanie = (
            f"OFFSET('SŁOWNIK'!${litera}$2,0,0,"
            f"MAX(1,COUNTA('SŁOWNIK'!${litera}:${litera})-1),1)"
        )
        skoroszyt.defined_names.add(DefinedName(f"sl_{klucz}", attr_text=odwolanie))

    for nazwa, klucz in uklad.NAZWY_PRZELICZNIKOW.items():
        numer_wiersza = wiersze_ustawien[klucz]
        skoroszyt.defined_names.add(
            DefinedName(nazwa, attr_text=f"'USTAWIENIA'!$C${numer_wiersza}")
        )


# ----------------------------------------------------------------------
#  Punkt wejścia
# ----------------------------------------------------------------------


def zbuduj_skoroszyt(
    operacje: list[Operacja],
    slowniki: Slowniki,
    ustawienia: Ustawienia,
    zdarzenia: list[ZdarzenieAudytu] | None = None,
) -> Workbook:
    """Składa kompletny skoroszyt z danych aplikacji."""
    skoroszyt = Workbook()
    skoroszyt.remove(skoroszyt.active)
    for nazwa in ARKUSZE:
        skoroszyt.create_sheet(nazwa)

    _zbuduj_start(skoroszyt["START"], ustawienia)
    _zbuduj_formularz(skoroszyt["FORMULARZ"])
    _zbuduj_dane(skoroszyt["Dane"], operacje)
    _zbuduj_magazyn(skoroszyt["MAGAZYN"])
    _zbuduj_raporty(skoroszyt["RAPORTY"])
    _zbuduj_slownik(skoroszyt["SŁOWNIK"], slowniki)
    _zbuduj_historie(skoroszyt["HISTORIA"], zdarzenia or [])
    wiersze_ustawien = _zbuduj_ustawienia(skoroszyt["USTAWIENIA"], ustawienia)

    _dodaj_nazwy(skoroszyt, wiersze_ustawien)

    for nazwa, code_name in ARKUSZE.items():
        skoroszyt[nazwa].sheet_properties.codeName = code_name

    skoroszyt.active = skoroszyt.index(skoroszyt["START"])
    return skoroszyt
